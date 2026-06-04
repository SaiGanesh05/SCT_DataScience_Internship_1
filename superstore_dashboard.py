import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule

# -----------------------------
# CREATE SAMPLE SUPERSTORE DATA
# -----------------------------
np.random.seed(42)

n = 500

df = pd.DataFrame({
    "Order Date": pd.date_range("2024-01-01", periods=n, freq="D"),
    "Category": np.random.choice(
        ["Technology", "Furniture", "Office Supplies"], n),
    "Sales": np.random.randint(100, 5000, n),
    "Profit": np.random.randint(-500, 2000, n)
})

df["Month"] = df["Order Date"].dt.strftime("%Y-%m")

# -----------------------------
# ANALYSIS
# -----------------------------

sales_by_category = df.groupby("Category")["Sales"].sum().reset_index()

profit_by_category = df.groupby("Category")["Profit"].sum().reset_index()

sales_trend = df.groupby("Month")["Sales"].sum().reset_index()

# -----------------------------
# EXPORT TO EXCEL
# -----------------------------

file_name = "Superstore_Sales_Dashboard.xlsx"

with pd.ExcelWriter(file_name, engine="openpyxl") as writer:

    df.to_excel(writer,
                sheet_name="Raw_Data",
                index=False)

    sales_by_category.to_excel(
        writer,
        sheet_name="Sales_by_Category",
        index=False)

    profit_by_category.to_excel(
        writer,
        sheet_name="Profit_by_Category",
        index=False)

    sales_trend.to_excel(
        writer,
        sheet_name="Sales_Trend",
        index=False)

# -----------------------------
# OPEN WORKBOOK
# -----------------------------

wb = load_workbook(file_name)

# -----------------------------
# DASHBOARD SHEET
# -----------------------------

dashboard = wb.create_sheet("Dashboard")

dashboard["A1"] = "SUPERSTORE SALES DASHBOARD"
dashboard["A1"].font = Font(size=16, bold=True)

# KPIs
dashboard["A3"] = "Total Sales"
dashboard["B3"] = df["Sales"].sum()

dashboard["A4"] = "Total Profit"
dashboard["B4"] = df["Profit"].sum()

dashboard["A5"] = "Total Orders"
dashboard["B5"] = len(df)

# -----------------------------
# BAR CHART
# -----------------------------

ws1 = wb["Sales_by_Category"]

bar = BarChart()
bar.title = "Total Sales by Category"
bar.y_axis.title = "Sales"

data = Reference(ws1, min_col=2, min_row=1, max_row=4)
cats = Reference(ws1, min_col=1, min_row=2, max_row=4)

bar.add_data(data, titles_from_data=True)
bar.set_categories(cats)

dashboard.add_chart(bar, "D2")

# -----------------------------
# PIE CHART
# -----------------------------

ws2 = wb["Profit_by_Category"]

pie = PieChart()
pie.title = "Profit by Category"

labels = Reference(ws2, min_col=1, min_row=2, max_row=4)
data = Reference(ws2, min_col=2, min_row=1, max_row=4)

pie.add_data(data, titles_from_data=True)
pie.set_categories(labels)

dashboard.add_chart(pie, "D18")  

# -----------------------------
# LINE CHART
# -----------------------------

ws3 = wb["Sales_Trend"]

line = LineChart()
line.title = "Sales Trend Over Time"
line.y_axis.title = "Sales"

data = Reference(
    ws3,
    min_col=2,
    min_row=1,
    max_row=ws3.max_row
)

cats = Reference(
    ws3,
    min_col=1,
    min_row=2,
    max_row=ws3.max_row
)

line.add_data(data, titles_from_data=True)
line.set_categories(cats)

dashboard.add_chart(line, "L2")

# -----------------------------
# CONDITIONAL FORMATTING
# -----------------------------

raw = wb["Raw_Data"]

color_scale = ColorScaleRule(
    start_type="min",
    start_color="FF0000",
    mid_type="percentile",
    mid_value=50,
    mid_color="FFFF00",
    end_type="max",
    end_color="00FF00"
)

raw.conditional_formatting.add(
    f"D2:D{raw.max_row}",
    color_scale
)

# -----------------------------
# FORMAT HEADERS
# -----------------------------

for sheet in wb.sheetnames:
    ws = wb[sheet]

    for cell in ws[1]:
        cell.fill = PatternFill(
            start_color="1F4E78",
            end_color="1F4E78",
            fill_type="solid"
        )
        cell.font = Font(
            color="FFFFFF",
            bold=True
        )

# -----------------------------
# SAVE
# -----------------------------

wb.save(file_name)

print("="*50)
print("Dashboard Created Successfully!")
print("File:", file_name)
print("="*50)